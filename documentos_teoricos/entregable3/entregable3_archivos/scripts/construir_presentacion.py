"""Construye documentos_teoricos/entregable3/entregable3_presentacion.pptx (exposición del APF3).

Cubre solo lo que pidió el docente para esta exposición: Project Charter básico, diagrama de Gantt, diagramas de
procesos con la cobertura del alcance funcional y no funcional, diagramas lógico y físico de la base de datos y
capturas de la app. Usa la línea gráfica y las utilidades de la PPT del entregable 2. Uso:
    python diagramas_e3.py            # si cambian el Gantt o el esquema de la base
    python construir_presentacion.py
"""
import io
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches

AQUI = os.path.dirname(os.path.abspath(__file__))
E3 = os.path.dirname(AQUI)
TEORICOS = os.path.normpath(os.path.join(AQUI, "..", "..", ".."))
E2 = os.path.join(TEORICOS, "entregable2", "entregable2_archivos")
CAPTURAS = os.path.join(E3, "capturas")
SALIDA = os.path.join(TEORICOS, "entregable3", "entregable3_presentacion.pptx")
sys.path.insert(0, os.path.join(E2, "scripts"))

from construir_presentacion import (caja_texto, rect, tabla, tarjeta, ROJO, ROJO_UTP, AZUL,  # noqa: E402  (de la PPT del APF2)
                                    VERDE, GRIS, BLANCO, LOGO, EXPOSITORES)
from contenido_documento import TITULO  # noqa: E402
from corregir_apf2 import corte_titulo  # noqa: E402

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
VACIA = prs.slide_layouts[6]
numero = [0]


def diapositiva(seccion, titulo, notas):
    numero[0] += 1
    s = prs.slides.add_slide(VACIA)
    caja_texto(s, 0.55, 0.22, 8, 0.3, seccion.upper(), size=11, color=ROJO, bold=True)
    caja_texto(s, 0.55, 0.5, 10.6, 0.65, titulo, size=26, color=AZUL, bold=True)
    rect(s, 0.55, 1.15, 1.5, 0.05, ROJO)
    s.shapes.add_picture(LOGO, Inches(11.6), Inches(0.25), Inches(1.3), Inches(0.346))
    rect(s, 0, 7.1, 13.333, 0.4, ROJO)
    caja_texto(s, 0.4, 7.1, 8, 0.4, "Facturas S20  ·  Avance de Proyecto Final 3", size=11, color=BLANCO,
               anchor=MSO_ANCHOR.MIDDLE)
    caja_texto(s, 12.3, 7.1, 0.6, 0.4, str(numero[0]), size=11, color=BLANCO, anchor=MSO_ANCHOR.MIDDLE,
               align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notas
    return s


def imagen(slide, ruta, x, y, w, h, sin_titulo=False, marco=False):
    """Inserta una imagen ajustada a la caja (centrada). Con sin_titulo recorta el título dibujado dentro."""
    im = Image.open(ruta)
    if sin_titulo:
        corte = corte_titulo(im)
        if corte:
            im = im.crop((0, corte, im.width, im.height))
    datos = io.BytesIO()
    im.save(datos, "PNG")
    datos.seek(0)
    escala = min(w / im.width, h / im.height)
    ww, hh = im.width * escala, im.height * escala
    ox, oy = x + (w - ww) / 2, y + (h - hh) / 2
    if marco:
        rect(slide, ox - 0.04, oy - 0.04, ww + 0.08, hh + 0.08, BLANCO, line=GRIS)
    return slide.shapes.add_picture(datos, Inches(ox), Inches(oy), Inches(ww), Inches(hh))


def etiquetas(slide, x, y, w, titulo, items, color):
    """Tarjeta con la lista de requerimientos que cubre un diagrama."""
    tarjeta(slide, x, y, w, 0.55 + 0.3 * len(items), titulo, items, color=color, size=11.5)


# ---------------------------------------------------------------- portada
def portada():
    numero[0] += 1
    s = prs.slides.add_slide(VACIA)
    rect(s, 0, 0, 13.333, 7.5, ROJO_UTP)
    rect(s, 4.45, 0.55, 4.43, 1.28, BLANCO)
    s.shapes.add_picture(LOGO, Inches(4.95), Inches(0.8), Inches(3.43), Inches(0.913))
    caja_texto(s, 0.8, 2.0, 11.73, 0.6, "Facturas S20", size=34, color=BLANCO, bold=True, align=PP_ALIGN.CENTER)
    caja_texto(s, 1.2, 2.65, 10.93, 0.9, TITULO, size=20, color=BLANCO, align=PP_ALIGN.CENTER)
    rect(s, 6.07, 3.7, 1.2, 0.05, BLANCO)
    caja_texto(s, 1.0, 3.9, 11.33, 0.35, "Avance de Proyecto Final 3  ·  Curso Integrador I: Sistemas Software",
               size=15, color=BLANCO, align=PP_ALIGN.CENTER)
    caja_texto(s, 1.0, 4.3, 11.33, 0.3, "Sección 48935  ·  Docente: Ing. Carlos Alberto Effio Gonzáles", size=12,
               color=BLANCO, align=PP_ALIGN.CENTER)
    caja_texto(s, 3.5, 4.85, 6.33, 1.3, list(EXPOSITORES.values()), size=13, color=BLANCO, align=PP_ALIGN.CENTER)
    caja_texto(s, 1.0, 6.7, 11.33, 0.3, "Lima – Perú  ·  2026", size=11, color=BLANCO, align=PP_ALIGN.CENTER)
    s.notes_slide.notes_text_frame.text = (
        "Presentamos el tercer avance de Facturas S20: el acta del proyecto, el cronograma al corte de la semana 12, "
        "los procesos que cubren el alcance, la base de datos de la app y la app funcionando en el teléfono.")


# ---------------------------------------------------------------- charter
def charter():
    s = diapositiva("Gestión del proyecto", "Project Charter",
                    "Acta del proyecto en su versión básica. El objetivo es el que pidió el docente y coincide con el "
                    "título. El alcance cubre desde la captura de la factura hasta el reporte y la exportación; la "
                    "declaración ante la SUNAT queda fuera. El entrenamiento del modelo se hace en una PC del equipo "
                    "y la app corre sin servidor.")
    filas = [
        ["Campo", "Contenido"],
        ["Proyecto", "Facturas S20: " + TITULO],
        ["Patrocinador", "Ing. Carlos Alberto Effio Gonzáles (docente del curso)"],
        ["Líder y equipo", "Diego Damián Valdivia (líder) · Paulo Coronel Obregón · Morán Mateo Velásquez · "
                           "Fabrizzio Gonzales Maco"],
        ["Objetivo", "Desarrollar una aplicación móvil con IA multimodal (Gemma 4) para mejorar la exactitud de la "
                     "declaración mensual del NRUS en bodegas de S.J.L."],
        ["Alcance", "Captura y lectura de facturas en el teléfono, verificación, acumulado mensual, categoría y cuota, "
                    "vencimiento, reporte PDF, historial y exportación de datos"],
        ["Fuera del alcance", "Declarar ante la SUNAT, emitir comprobantes, contabilidad completa, iOS"],
        ["Entregables", "APF1 (sem. 4) · APF2 (sem. 8) · APF3 (sem. 12) · informe final y sustentación (sem. 18) · "
                        "app Android y modelo .litertlm"],
        ["Criterio de éxito", "Total y categoría iguales al cálculo de control; exactitud ≥ 90 % por campo; lectura "
                              "≤ 20 s en el teléfono de referencia"],
        ["Restricciones", "18 semanas · S/ 0 en licencias · al menos 50 % del código en Java · sin servidor"],
        ["Riesgos principales", "Ajuste fino sin mejora suficiente · pérdida de exactitud al convertir · lectura lenta "
                                "en gama media"],
    ]
    t = tabla(s, 0.55, 1.45, 12.2, filas, [2.2, 10], size=12, alto_fila=0.48).table
    for i in range(1, len(filas)):
        t.cell(i, 1).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT


# ---------------------------------------------------------------- gantt
def gantt():
    s = diapositiva("Gestión del proyecto", "Diagrama de Gantt al corte de la semana 12",
                    "Cronograma al corte de hoy, semana 12. Análisis, diseño, el front-end de las 20 pantallas y la "
                    "base de datos están completos. El reporte PDF y la exportación se adelantaron. Siguen en curso el "
                    "ajuste fino del modelo y su integración en el teléfono, y empezaron las pruebas.")
    imagen(s, os.path.join(E3, "gantt_semana12.png"), 0.45, 1.35, 12.45, 5.0)
    caja_texto(s, 0.55, 6.45, 12.2, 0.5,
               [[("Completo: ", True), ("análisis, diseño, pantallas P01–P20, base cifrada, reporte PDF, exportación "
                                        "y avisos.   ", False),
                 ("En curso: ", True), ("ajuste fino de Gemma 4 y su lectura en el teléfono; pruebas.", False)]],
               size=12.5, color=GRIS)


# ---------------------------------------------------------------- procesos
def proceso_tobe():
    s = diapositiva("Diagramas de procesos", "Proceso propuesto (TO-BE) con Facturas S20",
                    "El proceso propuesto cubre el ciclo completo del mes: la factura se registra al recibirla, la "
                    "app acumula el total, determina la categoría y la cuota y avisa el vencimiento; al cierre, el "
                    "titular comparte el reporte con su contador.")
    imagen(s, os.path.join(E2, "bpmn_02_tobe.png"), 0.45, 1.4, 9.1, 5.55, sin_titulo=True)
    etiquetas(s, 9.75, 1.45, 3.1, "Alcance funcional", ["RF-01 a RF-08 registro", "RF-09 a RF-14 NRUS",
                                                         "RF-15 a RF-17 reporte y cierre", "RF-21 exportación"], VERDE)
    etiquetas(s, 9.75, 3.55, 3.1, "Alcance no funcional", ["RNF-03 total exacto", "RNF-09 tres toques",
                                                            "RNF-12 sin conexión", "RNF-14 datos en el teléfono"], AZUL)


def proceso_registro():
    s = diapositiva("Diagramas de procesos", "Subproceso «Registrar factura con IA»",
                    "Es el caso de uso principal. Tiene sus excepciones: foto ilegible, modelo no instalado o sin "
                    "memoria (pasa al registro manual), campos de baja confianza y factura duplicada. Todo ocurre "
                    "dentro del teléfono y el guardado es una sola transacción.")
    imagen(s, os.path.join(E2, "bpmn_03_registro.png"), 0.45, 1.4, 9.1, 5.55, sin_titulo=True)
    etiquetas(s, 9.75, 1.45, 3.1, "Alcance funcional", ["RF-01 a RF-03 captura", "RF-04 lectura en el teléfono",
                                                         "RF-05 y RF-06 verificación", "RF-07 duplicados",
                                                         "RF-08 total del mes"], VERDE)
    etiquetas(s, 9.75, 3.85, 3.1, "Alcance no funcional", ["RNF-01, 02 exactitud", "RNF-04, 05 tiempo y memoria",
                                                            "RNF-09, 10 usabilidad", "RNF-12, 14 sin red",
                                                            "RNF-13 transacción"], AZUL)


def procesos_soporte():
    s = diapositiva("Diagramas de procesos", "Subprocesos de cierre mensual y de publicación del modelo",
                    "El cierre mensual cubre la parte del NRUS: ventas, categoría, cuota, avisos, vencimiento, reporte "
                    "y cierre del mes. La publicación del modelo cubre cómo llega el modelo ajustado al teléfono: se "
                    "evalúa contra el base, se convierte y la app lo descarga una sola vez y verifica su huella.")
    imagen(s, os.path.join(E2, "bpmn_04_cierre.png"), 0.45, 1.35, 9.1, 2.75, sin_titulo=True)
    imagen(s, os.path.join(E2, "bpmn_05_modelo.png"), 0.45, 4.2, 9.1, 2.75, sin_titulo=True)
    etiquetas(s, 9.75, 1.35, 3.1, "Cierre mensual", ["RF-09 a RF-17, RF-21, RF-22", "RNF-03 total exacto",
                                                      "RNF-08 reporte PDF", "RNF-17, 18 reglas del NRUS"], VERDE)
    etiquetas(s, 9.75, 4.2, 3.1, "Publicación del modelo", ["RF-20 instalar y actualizar", "RNF-06, 07 tamaño y equipo",
                                                             "RNF-16 descarga verificada", "RNF-19 cambiar el modelo"],
              AZUL)


def cobertura():
    s = diapositiva("Diagramas de procesos", "Cobertura del alcance funcional y no funcional",
                    "Esta matriz resume qué proceso cubre cada requerimiento. Los 22 funcionales y los 19 no "
                    "funcionales quedan cubiertos. El acceso con PIN o huella y la anulación son transversales: no "
                    "tienen diagrama propio y se ven en las pantallas P01, P02 y P12.")
    filas = [
        ["Proceso", "Requerimientos funcionales", "Requerimientos no funcionales"],
        ["Proceso propuesto (TO-BE)", "RF-01 a RF-17, RF-21", "RNF-03, RNF-09, RNF-12, RNF-14"],
        ["Registrar factura con IA", "RF-01 a RF-08", "RNF-01, 02, 04, 05, 09, 10, 12, 13, 14"],
        ["Cierre mensual y declaración", "RF-09 a RF-17, RF-21, RF-22", "RNF-03, 08, 13, 17, 18"],
        ["Publicar y actualizar el modelo", "RF-20", "RNF-01, 02, 06, 07, 16, 19"],
        ["Transversal: acceso y anulación (P01, P02, P12)", "RF-18, RF-19", "RNF-11, RNF-15"],
        ["Total cubierto", "22 de 22 (100 %)", "19 de 19 (100 %)"],
    ]
    tabla(s, 0.55, 1.5, 12.2, filas, [4, 4, 4.2], size=13, alto_fila=0.55, resaltar={6})
    caja_texto(s, 0.55, 5.6, 12.2, 1.2,
               [[("Funcional: ", True), ("lo que hace la app (registrar, calcular, avisar, reportar, exportar).", False)],
                [("No funcional: ", True), ("cómo lo hace: exactitud, tiempo, memoria, sin conexión, cifrado y "
                                            "reglas mantenibles.", False)]],
               size=13, color=GRIS)


# ---------------------------------------------------------------- base de datos
def bd_logico():
    s = diapositiva("Base de datos", "Modelo lógico de la base de datos",
                    "Modelo lógico: 13 entidades en tercera forma normal con sus claves y relaciones. El centro es el "
                    "periodo (mes), que contiene las facturas y tiene una determinación con la categoría vigente. Los "
                    "parámetros del NRUS están en tablas versionadas, no en el código.")
    imagen(s, os.path.join(E3, "bd_logico.png"), 0.45, 1.35, 12.45, 5.65)


def bd_fisico_registro():
    s = diapositiva("Base de datos", "Modelo físico (1/2): registro de facturas",
                    "Modelo físico tal como lo crea Room en SQLite: tipos, NOT NULL, claves foráneas e índices. Los "
                    "montos se guardan en céntimos como INTEGER para no perder precisión; en Java son BigDecimal. Las "
                    "fotos no van en la base: se guarda la ruta del archivo cifrado.")
    imagen(s, os.path.join(E3, "bd_fisico_registro.png"), 0.45, 1.4, 12.45, 4.85)
    caja_texto(s, 0.55, 6.35, 12.2, 0.6,
               "SQLite + Room 2.8 · base cifrada completa con SQLCipher 4 (AES-256); la clave la protege el Android "
               "Keystore · montos en céntimos (INTEGER) · fechas ISO-8601 (TEXT)", size=12.5, color=GRIS)


def bd_fisico_nrus():
    s = diapositiva("Base de datos", "Modelo físico (2/2): parámetros y determinación del NRUS",
                    "La segunda parte del modelo físico guarda los parámetros del NRUS versionados, el cronograma de "
                    "vencimientos y el resultado de cada mes. Ahora mostramos la base real en DBeaver: el archivo del "
                    "teléfono está cifrado y no se abre; la copia legible muestra las tablas con los datos.")
    imagen(s, os.path.join(E3, "bd_fisico_nrus.png"), 0.45, 1.4, 8.6, 5.55)
    tarjeta(s, 9.3, 1.45, 3.55, 2.45, "Revisión en DBeaver",
            ["facturas_cifrada.db: el archivo del teléfono; sin la clave no abre",
             "facturas_legible.db: copia exportada dentro de la app",
             "Se genera con copiar_base.ps1"], color=VERDE, size=11.5)
    tarjeta(s, 9.3, 4.1, 3.55, 2.85, "Seguridad",
            ["Base cifrada con SQLCipher", "Clave en el Android Keystore", "Fotos cifradas con AES-GCM",
             "Sin copias automáticas a la nube", "Acceso con PIN o huella"], color=AZUL, size=11.5)


# ---------------------------------------------------------------- capturas
def capturas(seccion, titulo, notas, pantallas):
    s = diapositiva(seccion, titulo, notas)
    n = len(pantallas)
    ancho = 12.2 / n
    for i, (archivo, leyenda) in enumerate(pantallas):
        x = 0.55 + i * ancho
        imagen(s, os.path.join(CAPTURAS, archivo), x + 0.12, 1.45, ancho - 0.24, 4.95, marco=True)
        caja_texto(s, x, 6.5, ancho, 0.5, leyenda, size=12, color=AZUL, bold=True, align=PP_ALIGN.CENTER)


def capturas_todas():
    capturas("La app en el teléfono", "Acceso y preparación del lector",
             "La app se configura con el RUC del negocio y un PIN, y se abre con PIN o huella. El modelo Gemma 4 se "
             "descarga una sola vez después de comprobar la memoria y el espacio; desde ahí lee sin internet.",
             [("p01_configuracion.png", "P01 Configuración"), ("p02_acceso.png", "P02 Acceso"),
              ("p03b_modelo_instalado.png", "P03 Modelo instalado"), ("p04_inicio.png", "P04 Inicio")])
    capturas("La app en el teléfono", "Registro de una factura con Gemma 4",
             "El registro principal: se toma la foto con el encuadre, Gemma 4 lee la factura dentro del teléfono, el "
             "usuario revisa los campos (el de baja confianza aparece en ámbar), la app avisa si es duplicada y la "
             "guarda.",
             [("p06_camara.png", "P06 Cámara"), ("p07b_lectura_gemma.png", "P07 Lectura"),
              ("p08_verificacion.png", "P08 Verificación"), ("p09_duplicado.png", "P09 Duplicado"),
              ("p10_guardada.png", "P10 Guardada")])
    capturas("La app en el teléfono", "Control de facturas y determinación del NRUS",
             "Con las facturas del mes la app suma las compras; el titular ingresa sus ventas y la app aplica la "
             "regla del mayor monto para darle la categoría, la cuota y la fecha de vencimiento según su RUC.",
             [("p11_facturas.png", "P11 Facturas"), ("p12b_anulacion.png", "P12 Anulación"),
              ("p13_ventas.png", "P13 Ventas"), ("p14_categoria.png", "P14 Categoría"),
              ("p15_vencimiento.png", "P15 Vencimiento")])
    capturas("La app en el teléfono", "Reporte, historial, ajustes y exportación",
             "Al final del mes la app genera el reporte en PDF, cierra el mes y lo guarda en el historial. Si no hay "
             "memoria para el modelo, se registra a mano. Los datos se exportan a un ZIP con CSV para revisarlos en "
             "una PC.",
             [("p16_reporte.png", "P16 Reporte"), ("p17b_cerrar_mes.png", "P17 Cierre"),
              ("p18_ajustes.png", "P18 Ajustes"), ("p19_registro_manual.png", "P19 Registro manual"),
              ("p20_exportar.png", "P20 Exportar")])


def main():
    portada()
    charter()
    gantt()
    proceso_tobe()
    proceso_registro()
    procesos_soporte()
    cobertura()
    bd_logico()
    bd_fisico_registro()
    bd_fisico_nrus()
    capturas_todas()
    prs.save(SALIDA)
    print("Guardado en", SALIDA, f"({numero[0]} diapositivas)")


if __name__ == "__main__":
    main()

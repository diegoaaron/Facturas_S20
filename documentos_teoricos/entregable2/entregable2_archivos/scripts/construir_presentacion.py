"""Construye documentos_teoricos/entregable2/entregable2_presentacion.pptx (sustentación del APF2, ≈ 10 minutos).

Sigue la línea gráfica de la sustentación del entregable 1 (portada roja UTP, títulos azules, barra roja
inferior con el expositor) y usa las imágenes de esta carpeta. Uso:  python construir_presentacion.py
"""
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from contenido_documento import TITULO

AQUI = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.dirname(AQUI)
SALIDA = os.path.join(os.path.dirname(IMG), "entregable2_presentacion.pptx")
LOGO = os.path.join(IMG, "fuentes", "utp", "logo_utp.png")

ROJO = RGBColor(0xC0, 0x00, 0x00)
ROJO_UTP = RGBColor(0x9E, 0x1B, 0x32)
AZUL = RGBColor(0x1F, 0x4E, 0x79)
AZUL_CLARO = RGBColor(0xDC, 0xE6, 0xF2)
VERDE = RGBColor(0x1E, 0x7B, 0x4F)
VERDE_CLARO = RGBColor(0xE3, 0xF1, 0xE8)
NARANJA = RGBColor(0xC5, 0x5A, 0x11)
GRIS = RGBColor(0x40, 0x40, 0x40)
TEXTO = RGBColor(0x26, 0x26, 0x26)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONDO_CLARO = RGBColor(0xF6, 0xF8, 0xFA)
ROSA = RGBColor(0xFB, 0xE9, 0xE9)

EXPOSITORES = {"diego": "Damián Valdivia, Diego Aarón", "paulo": "Coronel Obregón, Paulo",
               "moran": "Mateo Velásquez, Morán", "fabrizzio": "Gonzales Maco, Fabrizzio Jhoel"}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
VACIA = prs.slide_layouts[6]
numero = [0]


# ---------------------------------------------------------------- utilidades
def caja_texto(slide, x, y, w, h, texto, size=14, color=TEXTO, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
               italic=False, font="Calibri"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    lineas = texto if isinstance(texto, list) else [texto]
    for i, linea in enumerate(lineas):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        partes = linea if isinstance(linea, list) else [(linea, bold)]
        for t, b in partes:
            r = par.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), b, italic, font
            r.font.color.rgb = color
    return tb


def vinetas(slide, x, y, w, h, items, size=13, color=TEXTO, espacio=4):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, it in enumerate(items):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.space_after = Pt(espacio)
        pPr = par._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.2)))
        pPr.set("indent", str(-Inches(0.2)))
        bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
        pPr.append(bu)
        partes = it if isinstance(it, list) else [(it, False)]
        for t, b in partes:
            r = par.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.name = Pt(size), b, "Calibri"
            r.font.color.rgb = color
    return tb


def rect(slide, x, y, w, h, fill, line=None, forma=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def imagen(slide, archivo, x, y, w, h, alinear="centro"):
    ruta = os.path.join(IMG, archivo)
    iw, ih = Image.open(ruta).size
    escala = min(w / iw, h / ih)
    ww, hh = iw * escala, ih * escala
    ox = x + (w - ww) / 2 if alinear == "centro" else x
    oy = y if alinear != "medio" else y + (h - hh) / 2
    return slide.shapes.add_picture(ruta, Inches(ox), Inches(oy), Inches(ww), Inches(hh))


def tarjeta(slide, x, y, w, h, titulo, cuerpo, color=AZUL, size=12, fondo=BLANCO):
    rect(slide, x, y, w, h, fondo, line=RGBColor(0xC9, 0xD3, 0xDD))
    rect(slide, x, y, w, 0.07, color)
    caja_texto(slide, x + 0.15, y + 0.15, w - 0.3, 0.35, titulo, size=13, color=color, bold=True)
    if isinstance(cuerpo, list):
        vinetas(slide, x + 0.15, y + 0.5, w - 0.3, h - 0.6, cuerpo, size=size)
    else:
        caja_texto(slide, x + 0.15, y + 0.5, w - 0.3, h - 0.6, cuerpo, size=size)


def cifra(slide, x, y, w, valor, etiqueta, color=AZUL, size=34):
    caja_texto(slide, x, y, w, 0.7, valor, size=size, color=color, bold=True, align=PP_ALIGN.CENTER)
    caja_texto(slide, x, y + 0.68, w, 0.5, etiqueta, size=11.5, color=GRIS, align=PP_ALIGN.CENTER)


def tabla(slide, x, y, w, filas, anchos, size=11, alto_fila=0.3, resaltar=None):
    nf, nc = len(filas), len(filas[0])
    shp = slide.shapes.add_table(nf, nc, Inches(x), Inches(y), Inches(w), Inches(alto_fila * nf))
    t = shp.table
    total = sum(anchos)
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(w * a / total)
    for i, fila in enumerate(filas):
        t.rows[i].height = Inches(alto_fila)
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.text = ""
            par = c.text_frame.paragraphs[0]
            r = par.add_run()
            r.text = str(v)
            r.font.size, r.font.name = Pt(size), "Calibri"
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = AZUL
                r.font.color.rgb, r.font.bold = BLANCO, True
                par.alignment = PP_ALIGN.CENTER
            else:
                c.fill.fore_color.rgb = (VERDE_CLARO if resaltar and i in resaltar else
                                         (FONDO_CLARO if i % 2 == 0 else BLANCO))
                r.font.color.rgb = TEXTO
                r.font.bold = j == 0 or bool(resaltar and i in resaltar)
                if j > 0:
                    par.alignment = PP_ALIGN.CENTER
    return shp


def diapositiva(seccion, titulo, expositor, notas):
    numero[0] += 1
    s = prs.slides.add_slide(VACIA)
    caja_texto(s, 0.55, 0.22, 8, 0.3, seccion.upper(), size=11, color=ROJO, bold=True)
    caja_texto(s, 0.55, 0.5, 10.6, 0.65, titulo, size=28, color=AZUL, bold=True)
    rect(s, 0.55, 1.2, 1.5, 0.05, ROJO)
    s.shapes.add_picture(LOGO, Inches(11.6), Inches(0.25), Inches(1.3), Inches(0.346))
    rect(s, 0, 7.1, 13.333, 0.4, ROJO)
    caja_texto(s, 0.4, 7.1, 6, 0.4, EXPOSITORES[expositor], size=11, color=BLANCO, anchor=MSO_ANCHOR.MIDDLE)
    caja_texto(s, 12.3, 7.1, 0.6, 0.4, str(numero[0]), size=11, color=BLANCO, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notas
    return s


# ---------------------------------------------------------------- 1. portada
def portada():
    numero[0] += 1
    s = prs.slides.add_slide(VACIA)
    rect(s, 0, 0, 13.333, 7.5, ROJO_UTP)
    rect(s, 4.45, 0.55, 4.43, 1.28, BLANCO)
    s.shapes.add_picture(LOGO, Inches(4.95), Inches(0.8), Inches(3.43), Inches(0.913))
    caja_texto(s, 0.8, 2.0, 11.73, 0.6, "Facturas S20", size=34, color=BLANCO, bold=True, align=PP_ALIGN.CENTER)
    caja_texto(s, 1.2, 2.65, 10.93, 0.9, TITULO, size=20, color=BLANCO, align=PP_ALIGN.CENTER)
    rect(s, 6.07, 3.7, 1.2, 0.05, BLANCO)
    caja_texto(s, 1.0, 3.9, 11.33, 0.35, "Avance de Proyecto Final 2  ·  Curso Integrador I: Sistemas Software", size=15, color=BLANCO, align=PP_ALIGN.CENTER)
    caja_texto(s, 1.0, 4.3, 11.33, 0.3, "Sección 48935  ·  Docente: Ing. Carlos Alberto Effio Gonzáles", size=12, color=BLANCO, align=PP_ALIGN.CENTER)
    caja_texto(s, 3.5, 4.85, 6.33, 1.3, list(EXPOSITORES.values()), size=13, color=BLANCO, align=PP_ALIGN.CENTER)
    caja_texto(s, 1.0, 6.7, 11.33, 0.3, "Lima – Perú  ·  2026", size=11, color=BLANCO, align=PP_ALIGN.CENTER)
    s.notes_slide.notes_text_frame.text = (
        "Diego: Buenos días. Somos el equipo de Facturas S20. En el primer avance analizamos el problema y elegimos la "
        "solución; hoy presentamos su diseño: procesos BPMN, arquitectura, clases, base de datos, prototipo, validación, "
        "cronograma y presupuesto. Cada integrante expone una parte.")


# ---------------------------------------------------------------- contenido
def contexto():
    s = diapositiva("Análisis del contexto", "El problema afecta a toda la cadena del proceso", "diego",
                    "Diego: Recordamos el problema con el enfoque de procesos que pide este avance. La bodega del NRUS recibe "
                    "facturas que no registra; eso afecta cinco procesos, desde la recepción hasta el pago. En el caso de referencia "
                    "se perdieron 5 de 14 facturas, el cálculo toma 3 horas y el titular pagó S/ 30 de más en el mes, S/ 360 al año. "
                    "La solución se organiza en seis módulos que cubren esos procesos.")
    cifras = [("14", "facturas de compra al mes"), ("5 de 14", "facturas perdidas al cierre"), ("3 h", "de cálculo manual al mes"),
              ("S/ 360", "pagados de más al año")]
    for i, (v, e) in enumerate(cifras):
        rect(s, 0.55 + i * 3.1, 1.45, 2.9, 1.25, FONDO_CLARO)
        cifra(s, 0.55 + i * 3.1, 1.5, 2.9, v, e, color=ROJO if i else AZUL, size=30)
    filas = [["Proceso de la bodega", "Cómo lo afecta el problema", "Qué hace Facturas S20"],
             ["Recepción de mercadería", "La factura se guarda sin registrar", "Foto en ≈ 10 s en el mostrador"],
             ["Custodia de comprobantes", "Se pierden o deterioran", "Imagen cifrada en el teléfono"],
             ["Cálculo mensual", "Suma a mano, 3 h y errores", "Acumulado automático"],
             ["Determinación de la categoría", "Ante la duda, categoría superior", "Motor de reglas y avisos al 80 %"],
             ["Declaración y pago", "Al último día y sin sustento", "Recordatorio y reporte PDF"]]
    tabla(s, 0.55, 3.0, 8.3, filas, [2.6, 2.9, 2.8], size=12, alto_fila=0.52)
    tarjeta(s, 9.1, 3.0, 3.7, 3.12, "Alcance: 6 módulos", ["Acceso y configuración", "Registro de facturas", "Control de facturas",
                                                          "Determinación del NRUS", "Reportes e historial", "Exportación y modelo"],
            color=VERDE, size=12.5)


def alternativas():
    s = diapositiva("Alternativas de solución", "Tres alternativas profundizadas; se mantiene la 1", "diego",
                    "Diego: Profundizamos las tres alternativas con los seis aspectos que pide la consigna. Las tres tienen al menos "
                    "cinco pantallas y más de 50 % de Java. La diferencia es dónde se lee la factura. Solo la alternativa 1 cubre "
                    "los 22 requerimientos funcionales y los 19 no funcionales, porque funciona sin conexión y la imagen no sale "
                    "del teléfono. Obtuvo 4,40 sobre 5 en la matriz ponderada. Paso la palabra a Paulo.")
    datos = [("Alternativa 1 · seleccionada", "App Android en Java con Gemma 4 embebido", VERDE,
              ["20 pantallas móviles", "Java ≈ 85 %", "Lee la factura sin internet", "La imagen no sale del teléfono", "RF 22/22 · RNF 19/19"], "4,40"),
             ("Alternativa 2", "App Android en Java con inferencia en servidor", NARANJA,
              ["7 pantallas", "Java ≈ 85 %", "Sin señal solo encola facturas", "Costo mensual de GPU", "RF 22/22 · RNF 16/19"], "3,60"),
             ("Alternativa 3", "Aplicación web en Java con servicio de terceros", GRIS,
              ["6 pantallas", "Java ≈ 80 %", "Modelo cerrado, sin ajuste fino", "Cobra por documento", "RF parcial · RNF 13/19"], "2,15")]
    for i, (a, b, col, items, pts) in enumerate(datos):
        x = 0.55 + i * 4.15
        rect(s, x, 1.5, 3.95, 0.95, col)
        caja_texto(s, x + 0.15, 1.55, 3.7, 0.3, a.upper(), size=11, color=BLANCO, bold=True)
        caja_texto(s, x + 0.15, 1.85, 3.7, 0.6, b, size=14, color=BLANCO, bold=True)
        rect(s, x, 2.45, 3.95, 3.3, BLANCO, line=RGBColor(0xC9, 0xD3, 0xDD))
        vinetas(s, x + 0.2, 2.6, 3.6, 2.2, items, size=13.5)
        caja_texto(s, x + 0.2, 4.85, 3.6, 0.8, [[("Puntaje ponderado  ", False), (pts, True)]], size=20, color=col)
    rect(s, 0.55, 5.95, 12.25, 0.95, VERDE_CLARO)
    caja_texto(s, 0.75, 6.0, 11.9, 0.85, "Criterios de mayor peso: permitir el ajuste fino sobre facturas peruanas (25 %) y funcionar sin "
               "conexión (20 %). La única desventaja de la alternativa 1 —requiere un teléfono de gama media— se mitiga con el "
               "registro manual (P19).", size=13, color=VERDE, anchor=MSO_ANCHOR.MIDDLE)


def asis():
    s = diapositiva("Diseño de procesos · BPMN", "Proceso actual (AS-IS): tres puntos de error", "paulo",
                    "Paulo: Modelamos los procesos con BPMN 2.0. Uso la notación estándar: pools por participante, eventos de "
                    "inicio, de temporizador y de fin, tareas y compuertas exclusivas. En el proceso actual, al fin de mes el titular "
                    "reúne, estima y suma a mano. La compuerta ¿Seguro de la categoría? es donde nace el pago en exceso. Las "
                    "anotaciones rojas son los tres puntos de error.")
    imagen(s, "bpmn_01_asis.png", 0.45, 1.4, 9.0, 5.5)
    tarjeta(s, 9.7, 1.45, 3.1, 2.55, "Notación usada", ["Pools: distribuidor, bodega, SUNAT", "Evento de temporizador: fin de mes",
                                                       "Compuertas exclusivas (XOR)", "Flujo de mensaje entre pools", "Almacén de datos: la caja"],
            size=11.5)
    tarjeta(s, 9.7, 4.15, 3.1, 2.75, "Puntos de error", ["Guardar sin registrar: se pierden 5 de 14", "Sumar a mano: 3 h y errores",
                                                        "Estimar ante la duda: paga S/ 50 en vez de S/ 20"], color=ROJO, size=12, fondo=ROSA)


def tobe():
    s = diapositiva("Diseño de procesos · BPMN", "Proceso propuesto (TO-BE): registrar al recibir", "paulo",
                    "Paulo: En el proceso propuesto la bodega tiene dos carriles: el bodeguero y la app. La factura se fotografía al "
                    "llegar, Gemma 4 la lee en el teléfono, el bodeguero solo verifica y el acumulado y la categoría se recalculan con "
                    "cada factura. Si se pasa del 80 % del límite, la app avisa. El cierre ya no es manual: un temporizador recuerda "
                    "el vencimiento y solo se pide el total de ventas.")
    imagen(s, "bpmn_02_tobe.png", 0.45, 1.4, 9.0, 5.55)
    filas = [["Aspecto", "AS-IS", "TO-BE"], ["Registro", "Fin de mes", "Al recibir"], ["Tiempo", "≈ 3 h", "≈ 10 s por factura"],
             ["Categoría", "Juicio ante la duda", "Motor de reglas"], ["Sustento", "64 % conservadas", "100 % con imagen"],
             ["Avisos", "Ninguno", "80 %, 100 % y vencimiento"]]
    tabla(s, 9.65, 1.5, 3.2, filas, [1.0, 1.05, 1.15], size=10.5, alto_fila=0.62)


def registro():
    s = diapositiva("Diseño de procesos · BPMN", "Subproceso «Registrar factura con IA» y sus excepciones", "paulo",
                    "Paulo: Detallamos el caso de uso principal en cuatro carriles. Tres compuertas cubren las excepciones: foto "
                    "ilegible, equipo sin memoria o sin modelo, y factura duplicada. Solo después de la revisión del bodeguero se "
                    "guarda todo en una transacción y se recalcula el mes. Los subprocesos de cierre mensual y de publicación del "
                    "modelo están en el Anexo K. Paso la palabra a Morán.")
    imagen(s, "bpmn_03_registro.png", 0.45, 1.4, 8.9, 5.55)
    tarjeta(s, 9.6, 1.45, 3.2, 3.05, "Excepciones diseñadas", [[("¿Legible? ", True), ("no → otra toma", False)],
                                                              [("¿Modelo y memoria? ", True), ("no → registro manual", False)],
                                                              [("¿Duplicada? ", True), ("sí → aviso, no se registra", False)],
                                                              [("¿Corrige? ", True), ("sí → editar campo", False)]], size=12)
    tarjeta(s, 9.6, 4.65, 3.2, 2.25, "Requerimientos", ["RF-01 a RF-08", "RNF-04: lectura ≤ 20 s", "RNF-12: sin conexión",
                                                       "RNF-14: la foto no sale"], color=VERDE, size=12)


def arquitectura():
    s = diapositiva("Diseño detallado", "Arquitectura: cuatro capas dentro del teléfono", "moran",
                    "Morán: La arquitectura tiene cuatro capas en el teléfono: presentación, dominio, inferencia local y datos. La "
                    "regla es que el dominio no depende de nadie: la inferencia y los datos implementan sus interfaces. Por eso el "
                    "motor del NRUS se prueba sin emulador y el modelo se cambia sin tocar el dominio. No hay servidor propio: fuera "
                    "del teléfono solo están Google Play y el repositorio público del que se descarga una vez el modelo; el "
                    "entrenamiento ocurre fuera del producto en Colab.")
    imagen(s, "arquitectura_capas.png", 0.45, 1.4, 8.4, 5.55)
    tarjeta(s, 9.1, 1.45, 3.7, 2.7, "Decisiones clave", ["Dominio en Java puro, sin Android", "Kotlin solo en el puente con LiteRT-LM",
                                                        "Room + SQLCipher con clave en el Keystore", "Modelo como descarga separada verificada"],
            size=12)
    tarjeta(s, 9.1, 4.3, 3.7, 2.6, "Tecnologías", ["Java 17 · Android 8.0+ · CameraX", "LiteRT-LM · Gemma 4 E2B (.litertlm)",
                                                  "Room + SQLCipher (SQLite)", "Exportación CSV/ZIP · sin servidor"], color=VERDE, size=12)


def clases():
    s = diapositiva("Diseño detallado", "Diagrama de clases: dominio y diseño", "moran",
                    "Morán: El diagrama de clases tiene dos vistas. El dominio, con 14 clases y enumeraciones para los estados; "
                    "agregamos ParametrosNrus, para cambiar la norma sin tocar el código, y Aviso. La vista de diseño muestra cómo se "
                    "implementa el registro por capas con buenas prácticas: inversión de dependencias, responsabilidad única y "
                    "encapsulamiento, con BigDecimal para los montos.")
    imagen(s, "clases_dominio.png", 0.45, 1.4, 6.2, 4.4)
    imagen(s, "clases_diseno.png", 6.75, 1.4, 6.1, 4.4)
    rect(s, 0.55, 5.95, 12.25, 0.95, FONDO_CLARO)
    caja_texto(s, 0.75, 6.0, 11.9, 0.85, [[("Buenas prácticas: ", True), ("inversión de dependencias (el caso de uso depende de interfaces), "
               "responsabilidad única (ViewModel solo estado), encapsulamiento (atributos privados, BigDecimal y enumeraciones).", False)]],
               size=13, color=AZUL, anchor=MSO_ANCHOR.MIDDLE)


def base_datos():
    s = diapositiva("Diseño detallado", "Base de datos: una sola base SQLite cifrada en el teléfono", "moran",
                    "Morán: La base de datos es una sola y vive en el teléfono: SQLite con Room, cifrada con SQLCipher. Son 13 "
                    "tablas en tercera forma normal y es la única fuente de verdad. Respecto del primer avance retiramos el "
                    "servidor de respaldo: no hay cuentas, respaldos remotos ni tablas de administración. Si el bodeguero quiere "
                    "revisar sus datos en una PC, la app los exporta a archivos CSV. Aplicamos consultas parametrizadas, PIN con "
                    "hash y verificación SHA-256 del modelo. Paso la palabra a Fabrizzio.")
    imagen(s, "der_local.png", 0.45, 1.4, 8.3, 4.3)
    caja_texto(s, 0.45, 5.72, 8.3, 0.3, "SQLite · Room + SQLCipher · 13 tablas · 3FN", size=12, color=AZUL, bold=True, align=PP_ALIGN.CENTER)
    tarjeta(s, 9.0, 1.45, 3.8, 4.5, "Exportación para PC", ["ZIP generado a pedido del titular", "facturas.csv y periodos.csv",
                                                           "determinaciones.csv", "Imágenes y reporte PDF (opcional)",
                                                           "Se abre en Excel o LibreOffice", "La app no lo envía a ningún servidor"],
            color=VERDE, size=12)
    medidas = ["SQLCipher (AES-256)", "Clave en el Keystore", "PIN con hash y sal", "Consultas parametrizadas", "Sin servidor ni cuentas", "SHA-256 del modelo"]
    for i, m in enumerate(medidas):
        x = 0.55 + i * 2.06
        rect(s, x, 6.15, 1.95, 0.75, VERDE_CLARO)
        caja_texto(s, x + 0.05, 6.15, 1.85, 0.75, m, size=11.5, color=VERDE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def ux():
    s = diapositiva("Diseño de interfaces · UX/UI", "Sistema de diseño y navegación", "fabrizzio",
                    "Fabrizzio: Las interfaces se diseñaron desde el proceso TO-BE: cada tarea del bodeguero tiene una pantalla. El "
                    "sistema de diseño fija colores con significado, tipografía, componentes y reglas de UX: botones de 48 dp, texto "
                    "de datos de 16 sp y lenguaje sin términos contables. La navegación usa una barra inferior con un botón central "
                    "Escanear, así registrar una factura está siempre a un toque.")
    imagen(s, "ui_sistema_diseno.png", 0.45, 1.4, 6.3, 3.6)
    imagen(s, "ui_mapa_navegacion.png", 6.85, 1.4, 6.0, 3.6)
    reglas = [("≥ 48 dp", "objetivos táctiles"), ("≥ 16 sp", "texto de datos"), ("≤ 3", "toques para registrar"), ("4,5:1", "contraste mínimo")]
    for i, (v, e) in enumerate(reglas):
        rect(s, 0.55 + i * 3.1, 5.4, 2.9, 1.5, FONDO_CLARO)
        cifra(s, 0.55 + i * 3.1, 5.5, 2.9, v, e, color=VERDE, size=28)


def prototipo_flujo():
    s = diapositiva("Diseño del prototipo", "Flujo principal: una factura en tres toques", "fabrizzio",
                    "Fabrizzio: Este es el flujo principal del prototipo. Desde el inicio, que muestra compras, categoría, cuota y "
                    "vencimiento, el bodeguero toca Escanear, encuadra la factura con indicador de nitidez, revisa los campos —el "
                    "importe aparece en ámbar porque la confianza fue baja— y guarda. Son tres toques. El total del mes se actualiza "
                    "y avisa que pasó el 80 % del límite.")
    pantallas = [("pantalla_04_inicio_resumen_del.png", "P04 Inicio"), ("pantalla_06_camara_con_encuadre.png", "P06 Cámara"),
                 ("pantalla_08_verificacion_de_datos.png", "P08 Verificación"), ("pantalla_10_factura_guardada.png", "P10 Guardada")]
    for i, (f, t) in enumerate(pantallas):
        x = 0.7 + i * 3.1
        imagen(s, f, x, 1.4, 2.6, 5.0)
        caja_texto(s, x, 6.45, 2.6, 0.35, t, size=13, color=AZUL, bold=True, align=PP_ALIGN.CENTER)
        if i < 3:
            flecha = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 2.68), Inches(3.75), Inches(0.35), Inches(0.35))
            flecha.fill.solid(); flecha.fill.fore_color.rgb = VERDE; flecha.line.fill.background()


def prototipo_modulos():
    s = diapositiva("Diseño del prototipo", "Módulos del NRUS, reportes y exportación", "fabrizzio",
                    "Fabrizzio: El prototipo completo tiene 20 pantallas móviles en seis módulos y el reporte PDF. Aquí la determinación: la app explica que decide el mayor monto entre compras y ventas. El "
                    "vencimiento se calcula con el último dígito del RUC. El reporte PDF es lo que el bodeguero comparte con su "
                    "contador, y la pantalla Exportar datos genera un ZIP con los CSV, las imágenes y el reporte para revisarlos en una "
                    "PC; la app no envía nada a ningún servidor.")
    imagen(s, "pantalla_14_categoria_y_cuota.png", 0.5, 1.4, 2.35, 4.6)
    imagen(s, "pantalla_15_vencimiento_y_recordatorio.png", 2.95, 1.4, 2.35, 4.6)
    imagen(s, "reporte_mensual_pdf.png", 5.45, 1.4, 3.1, 4.6)
    imagen(s, "pantalla_20_exportar_datos.png", 9.3, 1.4, 2.35, 4.6)
    for x, w, t in ((0.5, 2.35, "P14 Categoría y cuota"), (2.95, 2.35, "P15 Vencimiento"), (5.45, 3.1, "Reporte mensual PDF"),
                    (9.3, 2.35, "P20 Exportar datos")):
        caja_texto(s, x, 6.05, w, 0.3, t, size=12, color=AZUL, bold=True, align=PP_ALIGN.CENTER)
    rect(s, 0.55, 6.45, 12.25, 0.5, VERDE_CLARO)
    caja_texto(s, 0.75, 6.45, 11.9, 0.5, "Herramienta: mockups vectoriales SVG editables en Figma o Penpot · 20 pantallas móviles + reporte",
               size=12.5, color=VERDE, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def validacion():
    s = diapositiva("Validación de la solución", "El diseño cubre el 100 % del alcance", "fabrizzio",
                    "Fabrizzio: Validamos el diseño con matrices de cobertura. Cada uno de los 22 requerimientos funcionales tiene "
                    "pantalla, caso de uso, clase y tabla; los 19 no funcionales tienen una decisión de diseño verificable. "
                    "Hicimos una evaluación heurística con las diez heurísticas de Nielsen y un recorrido de seis tareas, y "
                    "definimos la prueba con tres titulares para la iteración 3. Devuelvo la palabra a Diego.")
    stats = [("22 / 22", "requerimientos funcionales cubiertos", VERDE), ("19 / 19", "no funcionales con decisión de diseño", VERDE),
             ("10 / 10", "heurísticas de Nielsen revisadas", AZUL), ("6", "tareas en la prueba con 3 titulares", AZUL)]
    for i, (v, e, c) in enumerate(stats):
        rect(s, 0.55 + i * 3.1, 1.45, 2.9, 1.45, FONDO_CLARO)
        cifra(s, 0.55 + i * 3.1, 1.55, 2.9, v, e, color=c, size=32)
    filas = [["Escenario verificado en el prototipo", "Resultado esperado", "Pantallas"],
             ["Factura legible", "Campos leídos; total S/ 4 120,50", "P06–P10"],
             ["Importe con baja confianza", "Campo en ámbar para revisar", "P08"],
             ["Factura duplicada", "Aviso; no se registra", "P09"],
             ["Equipo de 4 GB", "Registro manual validado", "P03, P19"],
             ["Más del 80 % del límite", "Barra y aviso en ámbar", "P04, P10, P14"],
             ["Anular por error", "Se descuenta y recalcula", "P12"]]
    tabla(s, 0.55, 3.15, 8.2, filas, [3.0, 3.2, 1.8], size=12, alto_fila=0.47)
    tarjeta(s, 9.0, 3.15, 3.8, 3.3, "Metas de usabilidad (I3)", ["≥ 5 de 6 tareas sin ayuda", "≤ 3 toques para registrar",
                                                                "≤ 30 s por factura", "SUS ≥ 70 puntos", "0 errores al leer la categoría"],
            color=VERDE, size=12.5)


def cronograma():
    s = diapositiva("Cronograma y presupuesto", "Cronograma al corte de la semana 8 y presupuesto", "diego",
                    "Diego: Al corte de la semana 8 el análisis y el diseño están completos y no hay desviaciones: adelantamos el "
                    "front-end una semana. El ajuste fino con conversión inicia la semana 9. El proyecto necesita un desembolso de "
                    "unos S/ 242; el resto son aportes del equipo. No hay costo de operación mensual: no existe servidor y la "
                    "lectura se hace en el teléfono.")
    imagen(s, "gantt_actualizado.png", 0.45, 1.4, 8.3, 4.1)
    cifs = [("S/ 242", "desembolso del proyecto", ROJO), ("S/ 19 342", "incluyendo aportes del equipo", AZUL),
            ("S/ 0", "operación al mes", VERDE), ("S/ 0", "costo por factura leída", VERDE)]
    for i, (v, e, c) in enumerate(cifs):
        y = 1.45 + i * 1.03
        rect(s, 9.05, y, 3.75, 0.93, FONDO_CLARO)
        caja_texto(s, 9.15, y + 0.05, 1.9, 0.8, v, size=22, color=c, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        caja_texto(s, 11.0, y + 0.05, 1.75, 0.8, e, size=11.5, color=GRIS, anchor=MSO_ANCHOR.MIDDLE)
    filas = [["Semanas", "Siguiente entrega", "Contenido"], ["9–10", "Incremento 1", "Front-end navegable y motor NRUS probado"],
             ["11–12", "APF3", "Room cifrada, modelo convertido en el teléfono, prueba de usabilidad"],
             ["13–18", "Final", "Reportes, exportación de datos, validación en mes simulado"]]
    tabla(s, 0.55, 5.65, 8.2, filas, [1.0, 1.6, 5.6], size=11.5, alto_fila=0.34)


def cierre():
    s = diapositiva("Documentación técnica · cierre", "Documentación para desarrolladores y próximos pasos", "diego",
                    "Diego: Por último, entregamos la documentación técnica en Markdown, pensada para versionarse con el código: "
                    "arquitectura, contratos, esquema de datos, reglas del NRUS, pantallas, exportación de datos y plan de pruebas. Con ella empieza la "
                    "construcción. En conclusión: el diseño cubre todo el alcance, protege los datos del bodeguero y mantiene 85 % "
                    "de Java. Gracias, quedamos atentos a sus preguntas.")
    tarjeta(s, 0.55, 1.45, 6.0, 4.4, "entregable2_documentacion_tecnica.md", [
        "Arquitectura por capas y regla de dependencias", "Reglas del NRUS con casos de prueba", "Contrato del modelo: instrucción y JSON de salida",
        "Esquema SQLite de 13 tablas y migraciones", "Pantallas P01–P20, tokens y textos", "Exportación de datos, seguridad y entornos",
        "Plan de construcción por iteraciones y pruebas"], size=13)
    tarjeta(s, 6.8, 1.45, 6.0, 2.6, "Próximos pasos hacia APF3", ["Front-end y dominio desde la documentación técnica",
                                                                 "Ajuste fino y conversión a .litertlm", "Exactitud y tiempo en el teléfono de referencia",
                                                                 "Prueba de usabilidad con tres titulares"], color=VERDE, size=13)
    rect(s, 6.8, 4.2, 6.0, 1.65, AZUL)
    caja_texto(s, 7.0, 4.25, 5.6, 1.55, "Una foto al recibir la factura, la IA la lee en el teléfono y el bodeguero declara "
               "el NRUS con el monto exacto.", size=17, color=BLANCO, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 0.55, 6.05, 12.25, 0.85, FONDO_CLARO)
    caja_texto(s, 0.75, 6.05, 11.9, 0.85, "Anexos del informe: A–J evidencias del primer avance · K–T evidencias del segundo avance "
               "(BPMN, clases, DER, arquitectura, interfaces, prototipo, cobertura, validación y documentación).",
               size=12.5, color=GRIS, anchor=MSO_ANCHOR.MIDDLE)


def main():
    portada(); contexto(); alternativas(); asis(); tobe(); registro(); arquitectura(); clases(); base_datos()
    ux(); prototipo_flujo(); prototipo_modulos(); validacion(); cronograma(); cierre()
    prs.core_properties.title = "APF2 - Sustentación — Facturas S20"
    prs.save(SALIDA)
    print("  guardado", SALIDA, f"({numero[0]} diapositivas)")


if __name__ == "__main__":
    main()
